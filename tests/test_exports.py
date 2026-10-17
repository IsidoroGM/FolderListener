"""Integración TXT/XLSX, texto literal y límites Excel."""

from datetime import datetime
from pathlib import Path, PurePosixPath, PureWindowsPath

import pytest
from openpyxl import load_workbook

from folderlister.config import Session
from folderlister.controller import generate
from folderlister.export_txt import TxtExporter
from folderlister.export_xlsx import XlsxExporter
from folderlister.issues import IssueLog
from folderlister.models import ElementKind, Entry, Representation, Statistics
from folderlister.rendering import output_stem


@pytest.mark.parametrize("representation", list(Representation))
def test_integrated_outputs(tree, representation):
    result = generate(Session(root=tree, recursive=True, representation=representation))
    assert len(result.outputs) == 2
    text = result.outputs[0].read_text(encoding="utf-8")
    assert "FolderLister 1.0.0" in text and "profundo.txt" in text
    assert ".oculto" not in text
    if representation == Representation.TREE:
        assert "│   └── profundo.txt" in text
    book = load_workbook(result.outputs[1], read_only=True)
    try:
        assert book.sheetnames == ["Resumen", "Elementos"]
        rows = list(book["Elementos"].values)
        assert len(rows) == 8
        assert rows[3][0] == 3
        assert rows[3][2] == "profundo.txt"
    finally:
        book.close()


def test_empty_workbook(tmp_path):
    result = generate(Session(root=tmp_path))
    assert "No se encontraron elementos" in result.outputs[0].read_text(encoding="utf-8")
    with result.outputs[1].open("rb") as source:
        book = load_workbook(source, read_only=True)
        assert len(list(book["Elementos"].values)) == 1
        book.close()


def test_formula_injection_and_splitting(tmp_path):
    names = ["=1+1.txt", "+SUM(1).txt", "-1.txt", "@SUM.txt", "normal.txt"]
    for name in names:
        (tmp_path / name).touch()
    result = generate(Session(root=tmp_path), max_rows=3)
    book = load_workbook(result.outputs[1])
    try:
        assert book.sheetnames == ["Resumen", "Elementos", "Elementos_2", "Elementos_3"]
        actual = []
        for sheet in book.worksheets[1:]:
            assert sheet.max_row <= 3
            for row in sheet.iter_rows(min_row=2):
                actual.append(row[2].value)
                assert all(cell.data_type != "f" for cell in row)
        assert set(actual) == set(names)
    finally:
        book.close()


def test_context_excluded_from_excel(tree):
    result = generate(
        Session(
            root=tree, recursive=True, kind=ElementKind.FILES, representation=Representation.TREE
        )
    )
    assert "[contexto; no incluido]" in result.outputs[0].read_text(encoding="utf-8")
    book = load_workbook(result.outputs[1], read_only=True)
    assert all(row[1] == "archivo" for row in list(book["Elementos"].values)[1:])
    book.close()


def test_output_names():
    assert output_stem(PurePosixPath("/")) == "raiz"
    assert output_stem(PureWindowsPath("C:/")) == "C"
    assert "\\" not in output_stem(PureWindowsPath("//server/share/"))
    assert len(output_stem(PurePosixPath("/" + "ñ" * 200)).encode()) <= 180


def test_long_cell_and_illegal_xml_preserved(tmp_path):
    stats = Statistics()
    issues = IssueLog(tmp_path / "issues.log", stats)
    config = Session(root=tmp_path)
    now = datetime.now().astimezone()
    txt = TxtExporter(tmp_path / "out.txt", tmp_path, config, now)
    xlsx = XlsxExporter(tmp_path / "out.xlsx", issues, txt.preserve)
    name = "😀" * 17000 + "\x01tail"
    entry = Entry(tmp_path / name, Path(name), False, True, (True,))
    try:
        txt.append(entry)
        xlsx.append(entry, config)
        xlsx.finish(tmp_path, config, now, stats)
        txt.finish(stats)
    finally:
        xlsx.close()
        txt.close()
        issues.close()
    assert stats.issues >= 1
    assert name in (tmp_path / "out.txt").read_text(encoding="utf-8")
    book = load_workbook(tmp_path / "out.xlsx", read_only=True)
    row = list(book["Elementos"].values)[1]
    assert len(row[2].encode("utf-16-le")) <= 65534
    book.close()


def test_incremental_large_stream(tmp_path):
    # 20 000 entradas sintéticas, sin crear archivos ni guardar filas en una lista.
    stats = Statistics()
    issues = IssueLog(tmp_path / "log", stats)
    writer = XlsxExporter(tmp_path / "large.xlsx", issues, lambda *args: None, max_rows=5001)
    config = Session(root=tmp_path)
    try:
        for index in range(20_000):
            name = f"archivo-{index}.txt"
            writer.append(Entry(tmp_path / name, Path(name), False, True, (True,)), config)
        writer.finish(tmp_path, config, datetime.now().astimezone(), stats)
    finally:
        writer.close()
        issues.close()
    book = load_workbook(tmp_path / "large.xlsx", read_only=True)
    assert len(book.sheetnames) == 5
    assert sum(sum(1 for _ in sheet.values) - 1 for sheet in book.worksheets[1:]) == 20_000
    book.close()
