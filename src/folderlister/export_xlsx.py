"""Excel en write_only, texto literal y partición automática de hojas."""

from collections.abc import Callable
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from folderlister import __version__
from folderlister.config import Session
from folderlister.issues import IssueLog
from folderlister.models import Entry, Statistics
from folderlister.rendering import render

HEADERS = (
    "Nivel de profundidad",
    "Tipo",
    "Nombre",
    "Ruta relativa",
    "Ruta absoluta",
    "Representación seleccionada",
)


def xml_text(value: str) -> str:
    """Escapa caracteres que XML 1.0 no permite sin producir un XLSX corrupto."""
    return "".join(
        character
        if (
            character in "\t\n\r"
            or 0x20 <= ord(character) <= 0xD7FF
            or 0xE000 <= ord(character) <= 0xFFFD
            or 0x10000 <= ord(character) <= 0x10FFFF
        )
        else f"\\u{ord(character):04x}"
        for character in value
    )


class XlsxExporter:
    """Escribe filas a disco; el límite inyectable facilita pruebas pequeñas."""

    def __init__(
        self,
        path: Path,
        issues: IssueLog,
        preserve: Callable[[Path, str, str], None],
        max_rows: int = 1_048_576,
    ) -> None:
        if not 2 <= max_rows <= 1_048_576:
            raise ValueError("El límite de filas debe estar entre 2 y 1048576.")
        self.path = path
        self.issues = issues
        self.preserve = preserve
        self.max_rows = max_rows
        self.book = Workbook(write_only=True)
        self.summary = self.book.create_sheet("Resumen")
        self.summary.column_dimensions["A"].width = 32
        self.summary.column_dimensions["B"].width = 90
        self.sheet_number = 0
        self._new_sheet()

    def _cell(self, sheet, value: object, path: Path, column: str, header: bool = False):
        if isinstance(value, str):
            clean = xml_text(value)
            # Excel cuenta unidades UTF-16, no puntos de código de Python.
            encoded = clean.encode("utf-16-le")
            if len(encoded) > 32767 * 2:
                clean = encoded[: 32767 * 2].decode("utf-16-le", errors="ignore")
            if clean != value:
                self.issues.record(
                    path,
                    f"adaptar celda {column}",
                    "Valor recortado o escapado; original íntegro en TXT",
                )
                self.preserve(path, column, value)
            cell = WriteOnlyCell(sheet, value=clean)
            cell.data_type = "s"  # Nunca fórmula, ni siquiera si comienza con '='.
            cell.number_format = "@"
        else:
            cell = WriteOnlyCell(sheet, value=value)
        if header:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="234E70")
        return cell

    def _new_sheet(self) -> None:
        self.sheet_number += 1
        name = "Elementos" if self.sheet_number == 1 else f"Elementos_{self.sheet_number}"
        self.sheet = self.book.create_sheet(name)
        self.sheet.freeze_panes = "A2"
        for column, width in zip("ABCDEF", (22, 14, 36, 65, 85, 85), strict=True):
            self.sheet.column_dimensions[column].width = width
        self.sheet.append([self._cell(self.sheet, text, self.path, text, True) for text in HEADERS])
        self.rows = 1

    def append(self, entry: Entry, config: Session) -> None:
        """Emite solo entradas incluidas; las carpetas de contexto no son filas."""
        if not entry.included:
            return
        if self.rows == self.max_rows:
            self._new_sheet()
        values = (
            entry.depth,
            "carpeta" if entry.is_directory else "archivo",
            entry.path.name,
            str(entry.relative),
            str(entry.path),
            render(entry, config.representation),
        )
        self.sheet.append(
            [
                self._cell(self.sheet, value, entry.path, column)
                for column, value in zip(HEADERS, values, strict=True)
            ]
        )
        self.rows += 1

    def finish(self, root: Path, config: Session, started: datetime, stats: Statistics) -> None:
        """Completa el resumen al finalizar la exploración y guarda el ZIP Excel."""
        rows = [
            ("Aplicación", f"FolderLister {__version__}"),
            ("Directorio", str(root)),
            ("Fecha y hora", started.isoformat(timespec="seconds")),
            ("Exploración", "recursiva" if config.recursive else "primer nivel"),
            ("Elementos incluidos", config.kind.value),
            ("Representación", config.representation.value),
            ("Directorios explorados", stats.directories_scanned),
            ("Carpetas listadas", stats.directories_listed),
            ("Archivos listados", stats.files_listed),
            ("Duración (s; exploración y filas)", round(stats.duration, 3)),
        ]
        self.summary.append(
            [self._cell(self.summary, text, root, text, True) for text in ("Campo", "Valor")]
        )
        for label, value in rows:
            self.summary.append(
                [
                    self._cell(self.summary, label, root, label),
                    self._cell(self.summary, value, root, label),
                ]
            )
        self.summary.append(["Número de errores / advertencias", stats.issues])
        self.summary.append(
            ["Estado", "Parcial / con advertencias" if stats.issues else "Completo"]
        )
        self.book.save(self.path)

    def close(self) -> None:
        """Libera también XML temporales en exportaciones interrumpidas.

        openpyxl 3.1.5 no expone abort para write_only; este adaptador encapsula
        sus atributos internos, verificados con la versión fijada y pruebas.
        """
        for sheet in self.book.worksheets:
            if sheet._writer is not None:
                if not sheet.closed:
                    sheet.close()
                if Path(sheet._writer.out).exists():
                    sheet._writer.cleanup()
        self.book.close()
